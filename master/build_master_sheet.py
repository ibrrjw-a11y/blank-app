"""상품 마스터 시트(.xlsx) 생성기. 구글 드라이브에 업로드하면 구글시트로 변환된다."""

import re
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

FONT = "Arial"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
AUTO_FILL = PatternFill("solid", fgColor="EDEDED")
WARN_FILL = PatternFill("solid", fgColor="F8CBAD")
BORDER = Border(bottom=Side(style="thin", color="D9D9D9"))

ROWS = 500  # 수식·드롭다운이 미리 깔리는 행 수

PRICE_POLICIES = [
    "소비자가격",
    "상시할인가",
    "브랜드위크가",
    "원데이특가",
    "모바일라이브방송가",
    "공동구매가",
    "폐쇄몰가격",
    "오프라인가격",
]
STATUS = ["사용", "중지"]
SALE_STATUS = ["판매중", "품절", "판매중지"]
YN = ["Y", "N"]


def q(name):
    return f"'{name}'"


def bound(formula):
    """$A:$A 같은 열 전체 참조를 데이터 행 범위($A$2:$A$501)로 바꾼다."""
    return re.sub(r"\$([A-Z]+):\$([A-Z]+)\b", lambda m: f"${m[1]}$2:${m[2]}${ROWS + 1}", formula)


def lookup(key_cell, sheet, key_col, val_col):
    rng = f"{q(sheet)}!${key_col}:${key_col}"
    out = f"{q(sheet)}!${val_col}:${val_col}"
    return f'IFERROR(INDEX({out},MATCH({key_cell},{rng},0)),"")'


def build_sheet(wb, title, columns, data=None):
    """columns: (헤더, 너비, 종류, 값) 목록.
    종류: in(입력), auto(수식, 값=행번호 r을 받는 함수), code(텍스트 입력), list(드롭다운, 값=수식/목록)
    """
    ws = wb.create_sheet(title)
    for c, (name, width, kind, _) in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(c)].width = width
    ws.row_dimensions[1].height = 24
    ws.freeze_panes = "A2"
    last_col = get_column_letter(len(columns))
    ws.auto_filter.ref = f"A1:{last_col}{ROWS + 1}"

    for c, (name, width, kind, spec) in enumerate(columns, start=1):
        col = get_column_letter(c)
        for r in range(2, ROWS + 2):
            cell = ws.cell(row=r, column=c)
            if kind == "auto":
                cell.value = "=" + bound(spec(r))
                cell.fill = AUTO_FILL
            elif kind == "code":
                cell.number_format = "@"
        if kind == "list":
            src = spec if isinstance(spec, str) else '"' + ",".join(spec) + '"'
            dv = DataValidation(type="list", formula1=src, allow_blank=True,
                                showErrorMessage=True, errorTitle="입력 오류",
                                error="목록에 있는 값만 선택할 수 있습니다.")
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{ROWS + 1}")

    for r, row in enumerate(data or [], start=2):
        for c, v in enumerate(row, start=1):
            if v is not None:
                ws.cell(row=r, column=c, value=v)
    return ws


def number_format(ws, col, fmt):
    for r in range(2, ROWS + 2):
        ws[f"{col}{r}"].number_format = fmt


def highlight(ws, ref, formula):
    ws.conditional_formatting.add(ref, FormulaRule(formula=[bound(formula)], fill=WARN_FILL))


WON = "#,##0"
PCT = "0.0%"


def main(path):
    wb = Workbook()
    wb.remove(wb.active)

    # 1. SKU
    sku = build_sheet(wb, "SKU", [
        ("SKU코드", 16, "code", None),
        ("상품명", 40, "in", None),
        ("브랜드", 16, "in", None),
        ("규격", 16, "in", None),
        ("원가", 12, "in", None),
        ("바코드", 16, "code", None),
        ("사방넷품번코드", 16, "code", None),
        ("사방넷단품코드", 16, "code", None),
        ("상태", 10, "list", STATUS),
    ])
    number_format(sku, "E", WON)
    highlight(sku, f"A2:A{ROWS + 1}", 'AND($A2<>"",COUNTIF($A:$A,$A2)>1)')

    # 2. 판매채널
    ch = build_sheet(wb, "판매채널", [
        ("채널코드", 14, "code", None),
        ("채널명", 20, "in", None),
        ("사방넷연동", 12, "list", YN),
        ("송신권한", 12, "list", YN),
        ("쇼핑몰ID", 20, "code", None),
        ("수수료율", 12, "in", None),
        ("상태", 10, "list", STATUS),
    ], data=[
        ["CPG", "쿠팡", "Y", "N", None, None, "사용"],
        ["SST", "스마트스토어", "Y", "N", None, None, "사용"],
        ["SHOPL", "슈플", "N", "N", None, None, "사용"],
    ])
    number_format(ch, "F", PCT)
    highlight(ch, f"A2:A{ROWS + 1}", 'AND($A2<>"",COUNTIF($A:$A,$A2)>1)')

    # 3. 광고채널
    ad = build_sheet(wb, "광고채널", [
        ("광고코드", 14, "code", None),
        ("매체명", 20, "in", None),
        ("유형", 14, "list", ["검색", "SNS", "디스플레이", "인플루언서", "공동구매", "라이브커머스", "기타"]),
        ("과금방식", 14, "list", ["CPC", "CPM", "CPA", "CPS", "고정비", "기타"]),
        ("상태", 10, "list", STATUS),
    ])
    highlight(ad, f"A2:A{ROWS + 1}", 'AND($A2<>"",COUNTIF($A:$A,$A2)>1)')

    # 4. 판매구성
    comp = build_sheet(wb, "판매구성", [
        ("구성코드", 18, "code", None),
        ("구성명", 40, "in", None),
        ("사방넷코드", 16, "code", None),
        ("구성원가", 12, "auto", lambda r: f'IF(A{r}="","",SUMIFS({q("구성품")}!$F:$F,{q("구성품")}!$A:$A,A{r}))'),
        ("구성품수", 10, "auto", lambda r: f'IF(A{r}="","",COUNTIF({q("구성품")}!$A:$A,A{r}))'),
        ("상태", 10, "list", STATUS),
    ])
    number_format(comp, "D", WON)
    highlight(comp, f"A2:A{ROWS + 1}", 'AND($A2<>"",COUNTIF($A:$A,$A2)>1)')
    highlight(comp, f"E2:E{ROWS + 1}", 'AND($A2<>"",$E2=0)')

    # 5. 구성품
    parts = build_sheet(wb, "구성품", [
        ("구성코드", 18, "list", f"={q('판매구성')}!$A$2:$A${ROWS + 1}"),
        ("구성명", 36, "auto", lambda r: f'IF(A{r}="","",{lookup(f"A{r}", "판매구성", "A", "B")})'),
        ("SKU코드", 16, "list", f"={q('SKU')}!$A$2:$A${ROWS + 1}"),
        ("상품명", 36, "auto", lambda r: f'IF(C{r}="","",{lookup(f"C{r}", "SKU", "A", "B")})'),
        ("수량", 8, "in", None),
        ("원가합", 12, "auto", lambda r: f'IF(OR(C{r}="",E{r}=""),"",E{r}*N({lookup(f"C{r}", "SKU", "A", "E")}))'),
    ])
    number_format(parts, "F", WON)

    # 6. 가격표: 구성 × 가격정책
    price_cols = [
        ("구성코드", 18, "list", f"={q('판매구성')}!$A$2:$A${ROWS + 1}"),
        ("구성명", 36, "auto", lambda r: f'IF(A{r}="","",{lookup(f"A{r}", "판매구성", "A", "B")})'),
        ("구성원가", 12, "auto", lambda r: f'IF(A{r}="","",{lookup(f"A{r}", "판매구성", "A", "D")})'),
    ] + [(p, 14, "in", None) for p in PRICE_POLICIES]
    price = build_sheet(wb, "가격표", price_cols)
    first_p = get_column_letter(4)
    last_p = get_column_letter(3 + len(PRICE_POLICIES))
    for c in range(3, 4 + len(PRICE_POLICIES)):
        number_format(price, get_column_letter(c), WON)
    price.freeze_panes = "D2"
    highlight(price, f"A2:A{ROWS + 1}", 'AND($A2<>"",COUNTIF($A:$A,$A2)>1)')
    # 판매가가 원가보다 낮으면 표시
    highlight(price, f"{first_p}2:{last_p}{ROWS + 1}", f'AND({first_p}2<>"",$C2<>"",{first_p}2<$C2)')

    # 7. 설정
    st = wb.create_sheet("설정")
    st["A1"], st["B1"] = "항목", "값"
    for cell in (st["A1"], st["B1"]):
        cell.font = Font(name=FONT, bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
    st["A2"], st["B2"] = "공급가 마진배분율", 0.5
    st["A2"].font = Font(name=FONT)
    st["B2"].font = Font(name=FONT, color="0000FF")
    st["B2"].number_format = "0%"
    st.column_dimensions["A"].width = 22
    st.column_dimensions["B"].width = 10

    # 8. 채널옵션: 구성 × 채널 × 가격정책 (+광고)
    price_block = f"{q('가격표')}!${first_p}:${last_p}"
    price_keys = f"{q('가격표')}!$A:$A"
    price_head = f"{q('가격표')}!${first_p}$1:${last_p}$1"

    def sell_price(r):
        cell = (f'INDEX({price_block},MATCH(C{r},{price_keys},0),'
                f'MATCH(E{r},{price_head},0))')
        return f'IF(OR(C{r}="",E{r}=""),"",IFERROR(IF({cell}="","",{cell}),""))'

    opt = build_sheet(wb, "채널옵션", [
        ("채널코드", 12, "list", f"={q('판매채널')}!$A$2:$A${ROWS + 1}"),
        ("채널명", 14, "auto", lambda r: f'IF(A{r}="","",{lookup(f"A{r}", "판매채널", "A", "B")})'),
        ("구성코드", 18, "list", f"={q('판매구성')}!$A$2:$A${ROWS + 1}"),
        ("구성명", 32, "auto", lambda r: f'IF(C{r}="","",{lookup(f"C{r}", "판매구성", "A", "B")})'),
        ("가격정책", 16, "list", f"={price_head}"),
        ("광고코드", 12, "list", f"={q('광고채널')}!$A$2:$A${ROWS + 1}"),
        ("쇼핑몰상품코드", 18, "code", None),
        ("쇼핑몰옵션명", 36, "in", None),
        ("사방넷코드", 14, "auto", lambda r: f'IF(C{r}="","",{lookup(f"C{r}", "판매구성", "A", "C")})'),
        ("판매가", 12, "auto", sell_price),
        ("구성원가", 12, "auto", lambda r: f'IF(C{r}="","",{lookup(f"C{r}", "판매구성", "A", "D")})'),
        ("공급가", 12, "auto", lambda r: f'IF(OR(J{r}="",K{r}=""),"",ROUND(J{r}-(J{r}-K{r})*{q("설정")}!$B$2,0))'),
        ("수수료", 12, "auto", lambda r: f'IF(OR(A{r}="",J{r}=""),"",ROUND(J{r}*N({lookup(f"A{r}", "판매채널", "A", "F")}),0))'),
        ("마진", 12, "auto", lambda r: f'IF(OR(J{r}="",K{r}=""),"",J{r}-M{r}-K{r})'),
        ("마진율", 10, "auto", lambda r: f'IF(OR(J{r}="",J{r}=0),"",N{r}/J{r})'),
        ("상태", 10, "list", SALE_STATUS),
    ])
    for col in "JKLMN":
        number_format(opt, col, WON)
    number_format(opt, "O", PCT)
    opt.freeze_panes = "E2"
    n = ROWS + 1
    # 가격표에 가격이 없는 조합
    highlight(opt, f"J2:J{n}", 'AND($C2<>"",$E2<>"",$J2="")')
    # 같은 채널에 같은 쇼핑몰상품코드+옵션명이 중복
    highlight(opt, f"G2:H{n}",
              'AND($G2<>"",COUNTIFS($A:$A,$A2,$G:$G,$G2,$H:$H,$H2)>1)')
    # 역마진
    highlight(opt, f"N2:N{n}", 'AND($N2<>"",$N2<0)')

    order = ["채널옵션", "가격표", "판매구성", "구성품", "SKU", "판매채널", "광고채널", "설정"]
    wb._sheets = [wb[name] for name in order]
    wb.active = 0
    wb.save(path)


if __name__ == "__main__":
    if len(sys.argv) > 2:
        ROWS = int(sys.argv[2])
    main(sys.argv[1] if len(sys.argv) > 1 else "상품마스터.xlsx")
